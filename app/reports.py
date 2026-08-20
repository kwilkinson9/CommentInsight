"""Builds a Word (.docx) report summarizing a document's comments -- meant
to be brought into a Comment Resolution Meeting (CRM).

Kept separate from storage.py (no DB access here, just dicts in -> bytes
out) so it can be tested and read on its own.
"""

from __future__ import annotations

import io
from datetime import datetime, timezone
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app import chart_images, storage

_MUTED = RGBColor(0x6B, 0x72, 0x80)
_BRAND_RED = RGBColor(0xB3, 0x18, 0x2F)
_BRAND_RED_HEX = "B3182F"
_BRAND_TINT_HEX = "FDEEF0"
_ICON_PATH = Path(__file__).resolve().parent / "static" / "comment-insight-icon.png"


def _add_bottom_border(paragraph, color: str = _BRAND_RED_HEX, size: int = 12) -> None:
    """A thin colored rule under a paragraph -- python-docx has no built-in
    horizontal rule, so this drops down to the underlying XML. Mirrors the
    red underline in the Comment Insight wordmark."""
    p_pr = paragraph._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), color)
    border.append(bottom)
    p_pr.append(border)


def _shade_cell(cell, color: str) -> None:
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), color)
    cell._tc.get_or_add_tcPr().append(shading)


def _add_section_heading(doc: Document, text: str):
    heading = doc.add_heading(text, level=1)
    for run in heading.runs:
        run.font.color.rgb = _BRAND_RED
    return heading


def _add_branded_header(doc: Document, title_text: str) -> None:
    """Icon + red title + underline rule -- the masthead every export shares."""
    if _ICON_PATH.exists():
        icon_paragraph = doc.add_paragraph()
        icon_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        icon_paragraph.add_run().add_picture(str(_ICON_PATH), width=Inches(0.55))

    title = doc.add_heading(title_text, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = _BRAND_RED
    _add_bottom_border(title)


def _add_credit_footer(doc: Document) -> None:
    doc.add_paragraph()
    credit = doc.add_paragraph()
    credit.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _muted(credit, "Comment Insight -- part of Dossentra, Medical Writing Solutions", italic=True)


def _friendly_date(value: str | None) -> str:
    if not value:
        return ""
    try:
        dt = datetime.fromisoformat(value.replace("Z", "+00:00")).astimezone()
    except ValueError:
        return value
    hour_12 = dt.hour % 12 or 12
    am_pm = "AM" if dt.hour < 12 else "PM"
    return f"{dt.strftime('%b')} {dt.day}, {dt.year}, {hour_12}:{dt.minute:02d} {am_pm}"


def _muted(paragraph, text: str, size: int = 9, italic: bool = False):
    run = paragraph.add_run(text)
    run.font.size = Pt(size)
    run.font.color.rgb = _MUTED
    run.italic = italic
    return run


def _add_comment_block(doc: Document, comment: dict) -> None:
    heading = doc.add_paragraph()
    heading.add_run(comment["author"] or "Unknown reviewer").bold = True
    _muted(heading, f"   {_friendly_date(comment['comment_date'])}")
    if comment.get("section"):
        _muted(heading, f"   {comment['section']}", italic=True)

    meta = doc.add_paragraph()
    category = comment.get("category") or "Not classified"
    status_label = storage.RESOLUTION_LABELS.get(comment.get("resolution_status"), "No decision yet")
    _muted(meta, f"Category: {category}    |    Resolution: {status_label}")
    if comment.get("resolution_note"):
        note_p = doc.add_paragraph()
        _muted(note_p, f"Note: {comment['resolution_note']}", italic=True)

    if comment.get("parent_author"):
        reply = doc.add_paragraph()
        _muted(reply, f"↳ reply to {comment['parent_author']}", italic=True)

    doc.add_paragraph(comment["text"])

    if comment.get("rationale"):
        rationale = doc.add_paragraph()
        _muted(rationale, comment["rationale"], italic=True)

    for conflict in comment.get("conflicts", []):
        conflict_p = doc.add_paragraph()
        conflict_p.add_run("Disagrees with another reviewer: ").bold = True
        conflict_p.add_run(f"{conflict['other_author']} — “{conflict['other_text']}”")
        reason_p = doc.add_paragraph()
        _muted(reason_p, conflict["reason"])

    if comment.get("anchor_text"):
        anchor = doc.add_paragraph()
        _muted(anchor, f"Document text: …{comment['anchor_text']}…")

    doc.add_paragraph()  # spacer between comments


def build_report(document: dict, comments: list[dict]) -> bytes:
    """Build a .docx report, grouped so comments needing team discussion
    (Decision Required, or flagged as a reviewer disagreement) come first,
    followed by a compact table of everything else."""
    doc = Document()
    _add_branded_header(doc, "Comment Resolution Report")

    filename_p = doc.add_paragraph(document["filename"])
    filename_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _muted(generated, f"Generated {_friendly_date(datetime.now(timezone.utc).isoformat())}")

    priority_comments = [c for c in comments if storage.is_priority(c)]
    other_comments = [c for c in comments if not storage.is_priority(c)]

    _add_section_heading(doc, "Summary")
    summary = doc.add_paragraph()
    summary.add_run(f"{len(comments)} total comment{'' if len(comments) == 1 else 's'}").bold = True
    summary.add_run(f"   ·   {len(priority_comments)} need team discussion")

    status_counts: dict[str, int] = {}
    for c in comments:
        status = c.get("resolution_status")
        if status:
            status_counts[status] = status_counts.get(status, 0) + 1
    if status_counts:
        parts = ", ".join(f"{storage.RESOLUTION_LABELS[s]}: {n}" for s, n in status_counts.items())
        doc.add_paragraph(f"Resolutions recorded so far: {parts}")

    if priority_comments:
        _add_section_heading(doc, "For Team Discussion")
        doc.add_paragraph(
            "Comments that need a specialist's sign-off, or where two reviewers were flagged as disagreeing."
        )
        for comment in priority_comments:
            _add_comment_block(doc, comment)

    if other_comments:
        _add_section_heading(doc, "Other Comments")
        table = doc.add_table(rows=1, cols=6)
        table.style = "Table Grid"
        headers = ["Reviewer", "Category", "Comment", "Document Text", "Resolution", "Note"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
            _shade_cell(cell, _BRAND_TINT_HEX)

        for comment in other_comments:
            row = table.add_row().cells
            row[0].text = comment["author"] or ""
            row[1].text = comment.get("category") or "Not classified"
            row[2].text = comment["text"]
            row[3].text = comment.get("anchor_text") or ""
            row[4].text = storage.RESOLUTION_LABELS.get(comment.get("resolution_status"), "No decision yet")
            row[5].text = comment.get("resolution_note") or ""

    _add_credit_footer(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


def build_multi_document_report(
    document_summaries: list[dict],
    chart_rows: list[dict],
    resolution_chart_rows: list[dict],
    total_documents: int,
    total_comments: int,
    priority_count: int,
    insights: dict | None = None,
    section_hotspot_rows: list[dict] | None = None,
) -> bytes:
    """Build the cross-document analysis report -- summary, AI insights,
    charts (embedded as images; python-docx can't render SVG), and a
    per-document breakdown table, in the same order the /analysis page
    leads with now (insights first, charts after). Unlike the
    single-document report, this doesn't re-list every comment -- that's
    what the Excel export and each document's own report are for."""
    doc = Document()
    _add_branded_header(doc, "Multi-Document Analysis Report")

    generated = doc.add_paragraph()
    generated.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _muted(generated, f"Generated {_friendly_date(datetime.now(timezone.utc).isoformat())}")

    _add_section_heading(doc, "Summary")
    summary = doc.add_paragraph()
    summary.add_run(f"{total_documents} document{'' if total_documents == 1 else 's'}").bold = True
    summary.add_run(f"   ·   {total_comments} total comment{'' if total_comments == 1 else 's'}")
    summary.add_run(f"   ·   {priority_count} need team discussion")

    if insights and (insights.get("overview") or insights.get("themes")):
        _add_section_heading(doc, "AI Insights")
        doc.add_paragraph(insights["overview"])
        for theme in insights.get("themes", []):
            theme_p = doc.add_paragraph()
            theme_p.add_run(theme["title"]).bold = True
            doc.add_paragraph(theme["description"])

    if chart_rows:
        _add_section_heading(doc, "Comments by Category")
        chart_png = chart_images.render_bar_chart_png("Comments by category", chart_rows)
        doc.add_picture(io.BytesIO(chart_png), width=Inches(6.3))

    if resolution_chart_rows:
        _add_section_heading(doc, "Comments by Resolution Status")
        chart_png = chart_images.render_bar_chart_png("Comments by resolution status", resolution_chart_rows)
        doc.add_picture(io.BytesIO(chart_png), width=Inches(6.3))

    if section_hotspot_rows:
        _add_section_heading(doc, "Section Hotspots (Recurring Across Documents)")
        doc.add_paragraph(
            "Sections that drew comments in more than one document -- often a sign the "
            "template's wording or instructions for that section are the recurring problem."
        )
        table = doc.add_table(rows=1, cols=3)
        table.style = "Table Grid"
        for cell, text in zip(table.rows[0].cells, ["Section", "Documents", "Comments"]):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
            _shade_cell(cell, _BRAND_TINT_HEX)
        for hotspot in section_hotspot_rows:
            row = table.add_row().cells
            row[0].text = hotspot["section"]
            row[1].text = str(hotspot["document_count"])
            row[2].text = str(hotspot["count"])

    if document_summaries:
        _add_section_heading(doc, "By Document")
        table = doc.add_table(rows=1, cols=7)
        table.style = "Table Grid"
        headers = ["Filename", "Comments", "Needs Discussion", "Accepted", "Rejected", "CRM", "No Decision"]
        for cell, text in zip(table.rows[0].cells, headers):
            cell.text = text
            cell.paragraphs[0].runs[0].bold = True
            _shade_cell(cell, _BRAND_TINT_HEX)

        for doc_summary in document_summaries:
            row = table.add_row().cells
            row[0].text = doc_summary["filename"]
            row[1].text = str(doc_summary["total"])
            row[2].text = str(doc_summary["priority_count"])
            row[3].text = str(doc_summary["accepted"])
            row[4].text = str(doc_summary["rejected"])
            row[5].text = str(doc_summary["crm"])
            row[6].text = str(doc_summary["no_decision"])

    _add_credit_footer(doc)

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
