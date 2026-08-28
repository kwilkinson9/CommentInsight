"""Tests for the Word (.docx) export report -- both the report-building
logic in app/reports.py directly, and the /documents/{id}/export route."""

import io
import pathlib
import sys
import unittest

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent.parent))

from docx import Document as DocxDocument

from app import reports
from tests.auth_helpers import AuthenticatedTestCase

SAMPLE = pathlib.Path(__file__).parent / "sample_docs" / "comment_insight_synthetic_sample.docx"


def _paragraph_texts(docx_bytes: bytes) -> list[str]:
    doc = DocxDocument(io.BytesIO(docx_bytes))
    return [p.text for p in doc.paragraphs]


class BuildReportTests(unittest.TestCase):
    """Exercises build_report() directly against plain dicts, no DB/HTTP."""

    def test_builds_a_valid_docx_with_expected_content(self):
        document = {"filename": "sample.docx"}
        comments = [
            {
                "id": 1, "author": "Priya Patel", "comment_date": "2026-01-05T10:00:00+00:00",
                "section": "Section 3", "category": "Decision Required", "resolution_status": None,
                "parent_author": None, "text": "We need to decide whether to list this as related.",
                "rationale": "Needs SME input.", "anchor_text": "the adverse event", "conflicts": [],
            },
            {
                "id": 2, "author": "Lisa Wong", "comment_date": "2026-01-06T10:00:00+00:00",
                "section": "Section 1", "category": "Editorial", "resolution_status": "accepted",
                "parent_author": None, "text": "Fix typo here.", "rationale": None,
                "anchor_text": "teh", "conflicts": [],
            },
        ]

        content = reports.build_report(document, comments)

        self.assertGreater(len(content), 0)
        # Must be a well-formed docx that python-docx can re-parse.
        text = "\n".join(_paragraph_texts(content))
        self.assertIn("Comment Resolution Report", text)
        self.assertIn("sample.docx", text)
        self.assertIn("We need to decide whether to list this as related.", text)
        self.assertIn("Priya Patel", text)

    def test_priority_comment_appears_in_discussion_section_not_table(self):
        document = {"filename": "sample.docx"}
        comments = [
            {
                "id": 1, "author": "Priya Patel", "comment_date": None, "section": None,
                "category": "Decision Required", "resolution_status": None, "parent_author": None,
                "text": "Needs team sign-off.", "rationale": None, "anchor_text": None, "conflicts": [],
            },
            {
                "id": 2, "author": "Lisa Wong", "comment_date": None, "section": None,
                "category": "Editorial", "resolution_status": None, "parent_author": None,
                "text": "Just a wording fix.", "rationale": None, "anchor_text": None, "conflicts": [],
            },
        ]

        content = reports.build_report(document, comments)
        text = "\n".join(_paragraph_texts(content))

        self.assertIn("For Team Discussion", text)
        self.assertIn("Other Comments", text)

        doc = DocxDocument(io.BytesIO(content))
        self.assertEqual(len(doc.tables), 1)
        table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn("Just a wording fix.", table_text)
        self.assertNotIn("Needs team sign-off.", table_text)

    def test_conflict_callout_included(self):
        document = {"filename": "sample.docx"}
        comments = [
            {
                "id": 1, "author": "Dr. Sarah Chen", "comment_date": None, "section": None,
                "category": "Scientific/Content", "resolution_status": None, "parent_author": None,
                "text": "This looks treatment-related to me.", "rationale": None, "anchor_text": None,
                "conflict_count": 1,
                "conflicts": [{"other_author": "James Okafor", "other_text": "Disagree, unrelated.", "reason": "They disagree about causality."}],
            },
        ]

        content = reports.build_report(document, comments)
        text = "\n".join(_paragraph_texts(content))

        self.assertIn("Disagrees with another reviewer", text)
        self.assertIn("James Okafor", text)
        self.assertIn("They disagree about causality.", text)

    def test_resolution_note_included_for_both_priority_and_table_comments(self):
        document = {"filename": "sample.docx"}
        comments = [
            {
                "id": 1, "author": "Priya Patel", "comment_date": None, "section": None,
                "category": "Decision Required", "resolution_status": "crm",
                "resolution_note": "Flagging for the CRM, needs biostats input.",
                "parent_author": None, "text": "Needs team sign-off.", "rationale": None,
                "anchor_text": None, "conflicts": [],
            },
            {
                "id": 2, "author": "Lisa Wong", "comment_date": None, "section": None,
                "category": "Editorial", "resolution_status": "accepted",
                "resolution_note": "Edited directly in the document.",
                "parent_author": None, "text": "Just a wording fix.", "rationale": None,
                "anchor_text": None, "conflicts": [],
            },
        ]

        content = reports.build_report(document, comments)
        text = "\n".join(_paragraph_texts(content))
        self.assertIn("Flagging for the CRM, needs biostats input.", text)

        doc = DocxDocument(io.BytesIO(content))
        table_text = " ".join(cell.text for row in doc.tables[0].rows for cell in row.cells)
        self.assertIn("Edited directly in the document.", table_text)

    def test_handles_no_comments(self):
        content = reports.build_report({"filename": "empty.docx"}, [])
        text = "\n".join(_paragraph_texts(content))
        self.assertIn("0 total comments", text)


class ExportRouteTests(AuthenticatedTestCase):
    def _upload_sample(self):
        with open(SAMPLE, "rb") as f:
            return self.client.post(
                "/upload",
                files={"file": ("comment_insight_synthetic_sample.docx", f, "application/vnd.openxmlformats-officedocument.wordprocessingml.document")},
            )

    def test_export_returns_a_valid_docx(self):
        self._upload_sample()

        resp = self.client.get("/documents/1/export")

        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp.headers["content-type"],
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        )
        self.assertIn("attachment", resp.headers["content-disposition"])
        self.assertIn(".docx", resp.headers["content-disposition"])

        text = "\n".join(_paragraph_texts(resp.content))
        self.assertIn("Comment Resolution Report", text)
        self.assertIn("comment_insight_synthetic_sample.docx", text)

    def test_export_missing_document_is_404(self):
        resp = self.client.get("/documents/999/export")
        self.assertEqual(resp.status_code, 404)


if __name__ == "__main__":
    unittest.main()
